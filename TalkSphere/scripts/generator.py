#!/usr/bin/env python3
"""
TalkSphere - Hebrew Word Database Generator
Main script for generating comprehensive Hebrew word content for VR/MR learning.

Usage:
    python generator.py                     # Generate all words
    python generator.py --test              # Test with first 3 words
    python generator.py --word אני          # Generate specific word
    python generator.py --start 10 --end 20 # Generate words 10-20
    python generator.py --resume            # Resume from checkpoint
"""

import argparse
import csv
import json
import logging
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Any

# Add scripts directory to path
sys.path.insert(0, str(Path(__file__).parent))

from api_handler import ClaudeAPIHandler, APIConfig, ContentGenerator
from validator import ContentValidator, validate_output_directory

# Configure logging
def setup_logging(log_dir: Path) -> logging.Logger:
    """Setup logging to both file and console."""
    log_dir.mkdir(parents=True, exist_ok=True)
    log_file = log_dir / f"generation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"

    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        handlers=[
            logging.FileHandler(log_file, encoding='utf-8'),
            logging.StreamHandler()
        ]
    )
    return logging.getLogger(__name__)


class TalkSphereGenerator:
    """Main generator orchestrating the content creation pipeline."""

    def __init__(self, base_path: str = None):
        """
        Initialize the generator.

        Args:
            base_path: Base path to TalkSphere directory. Defaults to parent of scripts dir.
        """
        if base_path:
            self.base_path = Path(base_path)
        else:
            self.base_path = Path(__file__).parent.parent

        self.input_path = self.base_path / "input"
        self.output_path = self.base_path / "output" / "words"
        self.templates_path = self.base_path / "templates" / "category_templates.json"
        self.logs_path = self.base_path / "logs"
        self.checkpoint_file = self.base_path / "checkpoint.json"

        self.logger = setup_logging(self.logs_path)
        self.api_handler = None
        self.validator = ContentValidator()
        self.templates = None

        # Statistics
        self.stats = {
            "start_time": None,
            "words_processed": 0,
            "words_successful": 0,
            "words_failed": 0,
            "categories_generated": 0,
            "total_items": 0,
            "errors": []
        }

    def load_templates(self):
        """Load category templates."""
        with open(self.templates_path, 'r', encoding='utf-8') as f:
            self.templates = json.load(f)
        self.validator.load_templates(str(self.templates_path))
        self.logger.info(f"Loaded {len(self.templates['categories'])} category templates")

    def load_word_list(self, file_path: str = None) -> List[Dict[str, Any]]:
        """
        Load word list from CSV or Excel file.

        Args:
            file_path: Path to word list file. If None, searches input directory.

        Returns:
            List of word dictionaries
        """
        if file_path:
            path = Path(file_path)
        else:
            # Search for word list in input directory
            csv_files = list(self.input_path.glob("*.csv"))
            xlsx_files = list(self.input_path.glob("*.xlsx"))

            if csv_files:
                path = csv_files[0]
            elif xlsx_files:
                path = xlsx_files[0]
            else:
                raise FileNotFoundError(f"No word list found in {self.input_path}")

        self.logger.info(f"Loading word list from: {path}")

        words = []

        if path.suffix == '.csv':
            with open(path, 'r', encoding='utf-8-sig') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    words.append(row)

        elif path.suffix == '.xlsx':
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path)
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[0]:  # Skip empty rows
                        words.append(dict(zip(headers, row)))
            except ImportError:
                self.logger.error("openpyxl not installed. Install with: pip install openpyxl")
                raise

        self.logger.info(f"Loaded {len(words)} words")
        return words

    def save_checkpoint(self, current_index: int, word_id: str):
        """Save progress checkpoint for resume capability."""
        checkpoint = {
            "last_index": current_index,
            "last_word": word_id,
            "timestamp": datetime.now().isoformat(),
            "stats": self.stats
        }
        with open(self.checkpoint_file, 'w', encoding='utf-8') as f:
            json.dump(checkpoint, f, ensure_ascii=False, indent=2)

    def load_checkpoint(self) -> Optional[int]:
        """Load checkpoint and return starting index."""
        if self.checkpoint_file.exists():
            with open(self.checkpoint_file, 'r', encoding='utf-8') as f:
                checkpoint = json.load(f)
            self.logger.info(f"Resuming from checkpoint: word {checkpoint['last_index']}")
            return checkpoint['last_index'] + 1
        return None

    def generate_word(self, word_data: Dict[str, Any], index: int, total: int) -> bool:
        """
        Generate all content for a single word.

        Args:
            word_data: Word information dictionary
            index: Current word index
            total: Total words count

        Returns:
            True if successful, False otherwise
        """
        word = word_data.get('word', '')
        word_nikud = word_data.get('word_nikud', word)
        transliteration = word_data.get('transliteration', '')
        translation = word_data.get('translation_en', '')
        word_id = word_data.get('id', str(index))

        # Create folder name
        folder_name = f"{int(word_id):03d}_{transliteration.lower().replace(' ', '_').replace('/', '_')}"
        word_output_dir = self.output_path / folder_name
        word_output_dir.mkdir(parents=True, exist_ok=True)

        self.logger.info(f"\n{'='*60}")
        self.logger.info(f"Word {index}/{total}: {word} ({transliteration}) - {translation}")
        self.logger.info(f"Output: {word_output_dir}")
        self.logger.info(f"{'='*60}")

        categories = self.templates['categories']
        successful_categories = 0
        word_items = 0

        for cat_idx, category in enumerate(categories, 1):
            self.logger.info(f"\n  Category {cat_idx}/24: {category['name_he']} ({category['name']})")

            output_file = word_output_dir / category['filename']

            # Skip if already generated
            if output_file.exists():
                self.logger.info(f"    Already exists, skipping...")
                successful_categories += 1
                continue

            try:
                # Generate content
                content = self.api_handler.generate_full_category(
                    word=word,
                    word_nikud=word_nikud,
                    transliteration=transliteration,
                    translation=translation,
                    category_info=category,
                    target_items=category['target_items']
                )

                if content:
                    # Add metadata
                    content['word'] = word
                    content['word_nikud'] = word_nikud
                    content['word_id'] = f"{transliteration.lower()}_{word_id}"
                    content['category'] = category['id']
                    content['generated_at'] = datetime.now().isoformat()

                    # Save to file
                    with open(output_file, 'w', encoding='utf-8') as f:
                        json.dump(content, f, ensure_ascii=False, indent=2)

                    # Validate
                    result = self.validator.validate_file(str(output_file))
                    items = result.total_items

                    self.logger.info(f"    ✓ Generated {items} items")
                    if result.warnings:
                        for w in result.warnings:
                            self.logger.warning(f"    Warning: {w}")

                    successful_categories += 1
                    word_items += items
                    self.stats['categories_generated'] += 1
                    self.stats['total_items'] += items

                else:
                    self.logger.error(f"    ✗ Failed to generate content")
                    self.stats['errors'].append({
                        "word": word,
                        "category": category['id'],
                        "error": "Generation returned None"
                    })

            except Exception as e:
                self.logger.error(f"    ✗ Error: {e}")
                self.stats['errors'].append({
                    "word": word,
                    "category": category['id'],
                    "error": str(e)
                })

        # Word complete
        success = successful_categories == len(categories)
        if success:
            self.stats['words_successful'] += 1
            self.logger.info(f"\n✓ Word complete: {word} - {successful_categories}/24 categories, {word_items:,} items")
        else:
            self.stats['words_failed'] += 1
            self.logger.warning(f"\n⚠ Word incomplete: {word} - {successful_categories}/24 categories")

        self.stats['words_processed'] += 1

        return success

    def run(
        self,
        start_index: int = 0,
        end_index: int = None,
        specific_word: str = None,
        resume: bool = False,
        test_mode: bool = False
    ):
        """
        Run the generation pipeline.

        Args:
            start_index: Starting word index (0-based)
            end_index: Ending word index (exclusive). None for all.
            specific_word: Generate only this specific word
            resume: Resume from checkpoint
            test_mode: Test with first 3 words only
        """
        self.stats['start_time'] = datetime.now().isoformat()

        # Load templates
        self.load_templates()

        # Initialize API handler
        self.api_handler = ClaudeAPIHandler()
        self.logger.info(f"API initialized: {self.api_handler.config.model}")

        # Load word list
        words = self.load_word_list()

        # Handle resume
        if resume:
            checkpoint_index = self.load_checkpoint()
            if checkpoint_index:
                start_index = checkpoint_index

        # Handle test mode
        if test_mode:
            self.logger.info("TEST MODE: Generating first 3 words only")
            words = words[:3]
            end_index = 3

        # Handle specific word
        if specific_word:
            words = [w for w in words if w.get('word') == specific_word]
            if not words:
                self.logger.error(f"Word not found: {specific_word}")
                return
            self.logger.info(f"Generating specific word: {specific_word}")

        # Handle range
        if end_index:
            words = words[start_index:end_index]
        elif start_index > 0:
            words = words[start_index:]

        total_words = len(words)
        self.logger.info(f"\n{'='*60}")
        self.logger.info(f"TALKSPHERE GENERATION STARTING")
        self.logger.info(f"Words to process: {total_words}")
        self.logger.info(f"Categories per word: 24")
        self.logger.info(f"Target items per category: ~2000")
        self.logger.info(f"Estimated total items: {total_words * 24 * 2000:,}")
        self.logger.info(f"{'='*60}\n")

        # Process words
        for idx, word_data in enumerate(words, 1):
            try:
                self.generate_word(word_data, idx, total_words)

                # Save checkpoint after each word
                self.save_checkpoint(start_index + idx - 1, word_data.get('id', str(idx)))

            except KeyboardInterrupt:
                self.logger.info("\n\nGeneration interrupted by user")
                self.save_checkpoint(start_index + idx - 1, word_data.get('id', str(idx)))
                break

            except Exception as e:
                self.logger.error(f"Critical error processing word: {e}")
                self.stats['errors'].append({
                    "word": word_data.get('word', 'unknown'),
                    "error": str(e)
                })
                continue

        # Generation complete
        self.generate_report()

    def generate_report(self):
        """Generate and save final report."""
        end_time = datetime.now()
        start_time = datetime.fromisoformat(self.stats['start_time']) if self.stats['start_time'] else end_time

        duration = end_time - start_time

        report = {
            "summary": {
                "start_time": self.stats['start_time'],
                "end_time": end_time.isoformat(),
                "duration_seconds": duration.total_seconds(),
                "words_processed": self.stats['words_processed'],
                "words_successful": self.stats['words_successful'],
                "words_failed": self.stats['words_failed'],
                "categories_generated": self.stats['categories_generated'],
                "total_items": self.stats['total_items'],
                "api_requests": self.api_handler.request_count if self.api_handler else 0
            },
            "errors": self.stats['errors']
        }

        # Save report
        report_path = self.logs_path / f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        with open(report_path, 'w', encoding='utf-8') as f:
            json.dump(report, f, ensure_ascii=False, indent=2)

        # Print summary
        self.logger.info(f"\n{'='*60}")
        self.logger.info("GENERATION COMPLETE")
        self.logger.info(f"{'='*60}")
        self.logger.info(f"Duration: {duration}")
        self.logger.info(f"Words processed: {self.stats['words_processed']}")
        self.logger.info(f"Words successful: {self.stats['words_successful']}")
        self.logger.info(f"Words failed: {self.stats['words_failed']}")
        self.logger.info(f"Categories generated: {self.stats['categories_generated']}")
        self.logger.info(f"Total items: {self.stats['total_items']:,}")
        self.logger.info(f"API requests: {self.api_handler.request_count if self.api_handler else 0}")
        self.logger.info(f"Report saved: {report_path}")

        if self.stats['errors']:
            self.logger.warning(f"\nErrors encountered: {len(self.stats['errors'])}")
            for error in self.stats['errors'][:10]:  # Show first 10
                self.logger.warning(f"  - {error.get('word', 'unknown')}: {error.get('error', 'unknown')}")


def main():
    """Main entry point."""
    parser = argparse.ArgumentParser(
        description="TalkSphere Hebrew Word Database Generator",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
    python generator.py --test              # Test with first 3 words
    python generator.py --word אני          # Generate specific word
    python generator.py --start 10 --end 20 # Generate words 10-20
    python generator.py --resume            # Resume from checkpoint
    python generator.py --validate          # Validate existing output
        """
    )

    parser.add_argument('--test', action='store_true',
                        help='Test mode: generate first 3 words only')
    parser.add_argument('--word', type=str,
                        help='Generate content for specific Hebrew word')
    parser.add_argument('--start', type=int, default=0,
                        help='Start index (0-based)')
    parser.add_argument('--end', type=int,
                        help='End index (exclusive)')
    parser.add_argument('--resume', action='store_true',
                        help='Resume from last checkpoint')
    parser.add_argument('--validate', action='store_true',
                        help='Validate existing output only')
    parser.add_argument('--base-path', type=str,
                        help='Base path to TalkSphere directory')

    args = parser.parse_args()

    # Initialize generator
    generator = TalkSphereGenerator(base_path=args.base_path)

    if args.validate:
        # Run validation only
        print("Running validation...")
        generator.load_templates()
        report = validate_output_directory(
            str(generator.output_path),
            str(generator.templates_path)
        )
        report_path = generator.logs_path / "validation_report.json"
        with open(report_path, 'w', encoding='utf-8') as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
        print(f"Validation report saved to: {report_path}")
        return

    # Run generation
    generator.run(
        start_index=args.start,
        end_index=args.end,
        specific_word=args.word,
        resume=args.resume,
        test_mode=args.test
    )


if __name__ == "__main__":
    main()
